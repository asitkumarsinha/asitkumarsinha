"""Command line entry point for the Hello world project."""

from pathlib import Path
from typing import Any, Dict, Iterable, List, Union


ExcelPath = Union[str, Path]
ExcelRow = Dict[str, Any]


def get_message() -> str:
    """Return the greeting displayed by the application."""
    return "Hello, world!"


def read_four_excel_files(file_paths: Iterable[ExcelPath]) -> Dict[str, List[ExcelRow]]:
    """Read exactly four Excel files into row dictionaries.

    The first row of each workbook's active sheet is treated as the header row.
    Empty rows are skipped, and blank headers are named ``column_<number>``.
    """
    paths = [Path(file_path) for file_path in file_paths]
    if len(paths) != 4:
        raise ValueError("Exactly four Excel file paths are required.")

    return {str(path): _read_excel_file(path) for path in paths}


def _read_excel_file(file_path: Path) -> List[ExcelRow]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            return []

        headers = [
            str(value) if value is not None else f"column_{index + 1}"
            for index, value in enumerate(header_row)
        ]

        records: List[ExcelRow] = []
        for row in rows:
            if all(value is None for value in row):
                continue

            record = {
                headers[index] if index < len(headers) else f"column_{index + 1}": value
                for index, value in enumerate(row)
            }
            records.append(record)

        return records
    finally:
        workbook.close()


def main() -> None:
    """Print the greeting to standard output."""
    print(get_message())


if __name__ == "__main__":
    main()
